"""Generación e importación de planillas XLSX bajo scripts/demo_planillas/<slug>/."""
from __future__ import annotations

import uuid
from datetime import timedelta
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text

from app.db import tenant_session
from app.modules.legajos.infra import ejecutar_comando
from app.modules.oc.lector_planilla_oc import leer_planilla_oc
from app.modules.operadoras.lector_xlsx import leer_planilla as leer_planilla_operadoras
from app.modules.operadoras import servicio as op_svc
from app.modules.oc import servicio as oc_svc
from app.auth.identidad import Rol
from scripts.demo.config import dni_tecnico, nombre_locacion_demo
from scripts.demo.contexto import EstadoTenant
from scripts.demo.db_util import ErrorDemo
from scripts.demo.fechas import hoy_tenant

RAIZ = Path(__file__).resolve().parents[1]
PLANTILLA = RAIZ.parents[1] / "frontend" / "public" / "Plantilla_presentaciones_operadoras.xlsx"
ENCABEZADOS_OC = [
    "Clave OC", "Referencia", "Operadora", "Locación", "Tipo de servicio",
    "Vigencia desde", "Vigencia hasta", "Estado",
]


def _dir_slug(est: EstadoTenant) -> Path:
    d = RAIZ / "demo_planillas" / est.spec.slug
    d.mkdir(parents=True, exist_ok=True)
    return d


def _guardar_oc(est: EstadoTenant, path: Path, filas_datos: list[list]) -> None:
    from html import escape
    from io import BytesIO
    from zipfile import ZIP_DEFLATED, ZipFile

    def celda(col: int, fila: int, valor: object) -> str:
        letras = ""
        n = col
        while n:
            n, resto = divmod(n - 1, 26)
            letras = chr(65 + resto) + letras
        return f'<c r="{letras}{fila}" t="inlineStr"><is><t>{escape(str(valor))}</t></is></c>'

    xml_filas = []
    for numero, valores in enumerate([ENCABEZADOS_OC] + filas_datos, start=1):
        xml_filas.append(f'<row r="{numero}">' + "".join(celda(i, numero, v) for i, v in enumerate(valores, 1)) + "</row>")
    contenido = BytesIO()
    with ZipFile(contenido, "w", ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
</Types>""")
        z.writestr("xl/workbook.xml", """<?xml version="1.0" encoding="UTF-8"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<sheets><sheet name="OC" sheetId="1" r:id="rId1"/></sheets></workbook>""")
        z.writestr("xl/_rels/workbook.xml.rels", """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
</Relationships>""")
        z.writestr(
            "xl/worksheets/sheet1.xml",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
            + "".join(xml_filas)
            + "</sheetData></worksheet>",
        )
    path.write_bytes(contenido.getvalue())


def _fechas_por_estado(hoy, estado: str) -> tuple[str, str, str]:
    """Columnas 11–13: exportación, presentación, respuesta (ISO con hora si aplica)."""
    dia = f"{hoy.isoformat()}T10:00:00+00:00"
    if estado == "exportado":
        return dia, "", ""
    if estado == "enviado":
        return dia, dia, ""
    if estado in ("aceptado", "rechazado"):
        return dia, dia, dia
    return "", "", ""


def generar_planillas(est: EstadoTenant) -> Path:
    out = _dir_slug(est)
    hoy = hoy_tenant(est.tenant_id)
    with tenant_session(est.tenant_id) as s:
        doc_id = s.execute(
            text(
                "SELECT d.documento_id::text FROM modulo1.documento d "
                "JOIN modulo1.definicion_requisito r "
                "ON r.tenant_id = d.tenant_id AND r.requisito_definicion_id = d.requisito_definicion_id "
                "WHERE d.tenant_id = :t AND d.sujeto_id = :s AND d.estado_version = 'vigente' "
                "AND r.nombre = 'Apto médico' LIMIT 1"
            ),
            {"t": est.tenant_id, "s": est.sujetos["tecnico1"]},
        ).scalar()
    assert doc_id, f"sin apto vigente para planilla operadoras ({est.spec.slug})"
    docs = [(op, str(doc_id)) for op in ("YPF", "Vista", "Tecpetrol")]
    wb = load_workbook(PLANTILLA)
    hoja = wb["Presentaciones"]

    def fila_presentacion(op: str, doc_id: str, estado: str, obs: str = "") -> None:
        nonlocal row
        f_exp, f_pres, f_resp = _fechas_por_estado(hoy, estado)
        hoja.cell(row=row, column=1, value=op)
        hoja.cell(row=row, column=2, value="persona")
        hoja.cell(row=row, column=3, value=dni_tecnico(est.spec.slug, 1))
        hoja.cell(row=row, column=5, value="Apto médico")
        hoja.cell(row=row, column=7, value=doc_id)
        hoja.cell(row=row, column=10, value=estado)
        if f_exp:
            hoja.cell(row=row, column=11, value=f_exp)
        if f_pres:
            hoja.cell(row=row, column=12, value=f_pres)
        if f_resp:
            hoja.cell(row=row, column=13, value=f_resp)
        hoja.cell(row=row, column=14, value=obs)
        row += 1

    def _escribir_fila_presentacion(
        hoja,
        row: int,
        *,
        operadora: str,
        doc_id: str,
        estado: str,
        slug: str,
        hoy,
        fecha_emision: str | None = None,
        fecha_vencimiento: str | None = None,
        requisito_id: str | None = None,
        observacion: str = "",
        incluir_vigencia: bool = False,
    ) -> None:
        f_exp, f_pres, f_resp = _fechas_por_estado(hoy, estado)
        hoja.cell(row=row, column=1, value=operadora)
        hoja.cell(row=row, column=2, value="persona")
        hoja.cell(row=row, column=3, value=dni_tecnico(slug, 1))
        if requisito_id is not None:
            hoja.cell(row=row, column=6, value=requisito_id)
        hoja.cell(row=row, column=5, value="Apto médico")
        if doc_id:
            hoja.cell(row=row, column=7, value=doc_id)
        if incluir_vigencia:
            emision = fecha_emision if fecha_emision is not None else hoy.isoformat()
            vencimiento = fecha_vencimiento if fecha_vencimiento is not None else (hoy + timedelta(days=365)).isoformat()
            hoja.cell(row=row, column=8, value=emision)
            hoja.cell(row=row, column=9, value=vencimiento)
        elif fecha_vencimiento is not None:
            hoja.cell(row=row, column=8, value=fecha_emision if fecha_emision is not None else hoy.isoformat())
            hoja.cell(row=row, column=9, value=fecha_vencimiento)
        hoja.cell(row=row, column=10, value=estado)
        if f_exp:
            hoja.cell(row=row, column=11, value=f_exp)
        if f_pres:
            hoja.cell(row=row, column=12, value=f_pres)
        if f_resp:
            hoja.cell(row=row, column=13, value=f_resp)
        if observacion:
            hoja.cell(row=row, column=14, value=observacion)

    row = 6
    for i, (op, doc) in enumerate(docs):
        fila_presentacion(op, doc, ("exportado", "enviado", "aceptado")[i])
    for op, doc in docs[:2]:
        fila_presentacion(op, doc, "rechazado", "Observación demo")
    p1 = out / "presentaciones_1.xlsx"
    wb.save(p1)

    wb2 = load_workbook(p1)
    hoja2 = wb2["Presentaciones"]
    doc_ypf = docs[0][1]
    dia2 = f"{(hoy + timedelta(days=1)).isoformat()}T11:00:00+00:00"
    dia3 = f"{(hoy + timedelta(days=2)).isoformat()}T11:00:00+00:00"
    row = hoja2.max_row + 1
    hoja2.cell(row=row, column=1, value="YPF")
    hoja2.cell(row=row, column=2, value="persona")
    hoja2.cell(row=row, column=3, value=dni_tecnico(est.spec.slug, 1))
    hoja2.cell(row=row, column=5, value="Apto médico")
    hoja2.cell(row=row, column=7, value=doc_ypf)
    hoja2.cell(row=row, column=10, value="enviado")
    hoja2.cell(row=row, column=11, value=dia2)
    hoja2.cell(row=row, column=12, value=dia2)
    row += 1
    hoja2.cell(row=row, column=1, value="YPF")
    hoja2.cell(row=row, column=2, value="persona")
    hoja2.cell(row=row, column=3, value=dni_tecnico(est.spec.slug, 1))
    hoja2.cell(row=row, column=5, value="Apto médico")
    hoja2.cell(row=row, column=7, value=doc_ypf)
    hoja2.cell(row=row, column=10, value="aceptado")
    hoja2.cell(row=row, column=11, value=dia2)
    hoja2.cell(row=row, column=12, value=dia3)
    hoja2.cell(row=row, column=13, value=dia3)
    wb2.save(out / "presentaciones_2.xlsx")

    wb_err = load_workbook(PLANTILLA)
    h_err = wb_err["Presentaciones"]
    doc_falso = "00000000-0000-0000-0000-000000000099"
    _escribir_fila_presentacion(
        h_err, 6, operadora="YPF SA", doc_id=doc_id, estado="exportado", slug=est.spec.slug, hoy=hoy
    )
    _escribir_fila_presentacion(
        h_err, 7, operadora="YPF", doc_id=doc_falso, estado="exportado", slug=est.spec.slug, hoy=hoy
    )
    _escribir_fila_presentacion(
        h_err,
        8,
        operadora="YPF",
        doc_id=doc_id,
        estado="exportado",
        slug=est.spec.slug,
        hoy=hoy,
        fecha_vencimiento="2099-99-99",
    )
    _escribir_fila_presentacion(
        h_err,
        9,
        operadora="YPF",
        doc_id=doc_id,
        estado="exportado",
        slug=est.spec.slug,
        hoy=hoy,
        requisito_id="NO-UUID",
    )
    _escribir_fila_presentacion(
        h_err, 10, operadora="", doc_id=doc_id, estado="exportado", slug=est.spec.slug, hoy=hoy
    )
    wb_err.save(out / "presentaciones_con_errores.xlsx")

    loc_ypf = nombre_locacion_demo("YPF", 1)
    oc_rows = [
        [
            f"OC-PLAN-{est.spec.slug}-OK",
            "ref",
            "YPF",
            loc_ypf,
            "Wireline",
            (hoy + timedelta(days=60)).isoformat(),
            (hoy + timedelta(days=90)).isoformat(),
            "activo",
        ],
        [
            f"OC-PLAN-{est.spec.slug}-BAD",
            "ref",
            "Operadora Inexistente",
            loc_ypf,
            "Wireline",
            hoy.isoformat(),
            (hoy + timedelta(days=10)).isoformat(),
            "activo",
        ],
    ]
    _guardar_oc(est, out / "oc.xlsx", oc_rows)
    return out


def _exigir_filas_importadas(etiqueta: str, resp: dict) -> None:
    if int(resp.get("filas_aceptadas") or 0) == 0:
        rech = resp.get("filas_rechazadas")
        errores = resp.get("errores") or []
        muestra = ""
        if errores:
            e0 = errores[0]
            muestra = f" — fila {e0.get('fila')}: {e0.get('codigo')} {e0.get('mensaje')}"
        raise ErrorDemo(f"{etiqueta}: importación sin filas aceptadas (rechazadas={rech}){muestra}")


def importar_planillas(est: EstadoTenant, idn) -> None:
    base = _dir_slug(est)
    for nombre in ("presentaciones_1.xlsx", "presentaciones_2.xlsx"):
        data = (base / nombre).read_bytes()
        filas, err = leer_planilla_operadoras(data)
        resp = ejecutar_comando(
            idn,
            str(uuid.uuid4()),
            (Rol.RESPONSABLE_LEGAJOS,),
            lambda s: op_svc.importar_filas(
                s, idn, archivo=nombre, hoja="Presentaciones", filas=filas, errores_lectura=err
            ),
            ruta="/comandos/importar_planilla_operadoras",
            body={"archivo": nombre},
        )
        _exigir_filas_importadas(f"planilla operadoras {nombre} ({est.spec.slug})", resp)
    lote = uuid.uuid5(uuid.NAMESPACE_DNS, f"import-oc-{est.spec.slug}")
    data_oc = (base / "oc.xlsx").read_bytes()
    filas, err = leer_planilla_oc(data_oc)
    with tenant_session(est.tenant_id) as s:
        from app.comun.idempotencia import ejecutar_idempotente, fingerprint_de

        resp = ejecutar_idempotente(
            est.tenant_id,
            idn.usuario_id,
            f"lote_oc:{lote}",
            fingerprint_de("POST", "/comandos/importar_planilla_oc", {"lote_id": str(lote)}),
            lambda s: oc_svc.importar_lote_oc(s, idn, str(lote), "planilla", filas),
        )
        _exigir_filas_importadas(f"planilla OC ({est.spec.slug})", resp)
