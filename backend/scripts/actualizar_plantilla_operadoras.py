"""Actualiza la plantilla pública de presentaciones de operadoras (columna Documento ID + Instrucciones)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[2]
DESTINOS = [
    RAIZ / "frontend" / "public" / "Plantilla_presentaciones_operadoras.xlsx",
]

ENCABEZADOS = [
    "Operadora", "Tipo de sujeto", "Identificador del sujeto", "Sujeto ID",
    "Tipo de documento", "Requisito ID", "Documento ID", "Fecha de emisión", "Fecha de vencimiento",
    "Estado", "Fecha de exportación", "Fecha de presentación", "Fecha de respuesta", "Observación",
]
COLUMNA_CONTROL = "Control"
FILA_ENCABEZADO = 5
FILA_DATOS_DESDE = 6
FILA_DATOS_HASTA = 50

INSTRUCCIONES = [
    "Complete la hoja Presentaciones. Cada fila actualiza el espejo de una operadora; no crea documentos del legajo.",
    "La operadora debe existir previamente en Catálogos OC (no se crean operadoras nuevas desde la planilla).",
    "Documento ID es opcional: use el UUID de la versión documental cuando haya ambigüedad.",
    "Fechas: celda de fecha Excel, ISO (aaaa-mm-dd) o texto dd/mm/aaaa.",
    "Estados válidos: exportado, enviado, aceptado, rechazado (con sus fechas asociadas).",
    "La columna Control es solo ayuda en Excel; no se importa al sistema.",
]

_FORMULA_CONTROL = (
    '=IF(COUNTA(A{row}:M{row})=0,"",'
    'IF(OR(A{row}="",B{row}="",C{row}="",E{row}="",J{row}=""),"Faltan campos obligatorios",'
    'IF(AND(J{row}="exportado",K{row}=""),"Falta fecha de exportación",'
    'IF(AND(OR(J{row}="enviado",J{row}="aceptado",J{row}="rechazado"),L{row}=""),"Falta fecha de presentación",'
    'IF(AND(OR(J{row}="aceptado",J{row}="rechazado"),M{row}=""),"Falta fecha de respuesta",'
    'IF(AND(I{row}<>"",H{row}<>"",I{row}<H{row}),"Vencimiento anterior a emisión","Lista para importar")))))'
)


def _formula_control(fila: int) -> str:
    return _FORMULA_CONTROL.format(row=fila)


def actualizar(ruta: Path) -> None:
    wb = load_workbook(ruta)
    if "Presentaciones" not in wb.sheetnames:
        raise SystemExit(f"Falta hoja Presentaciones en {ruta}")
    hoja = wb["Presentaciones"]
    for col, titulo in enumerate(ENCABEZADOS, start=1):
        hoja.cell(row=FILA_ENCABEZADO, column=col, value=titulo)
    hoja.cell(row=FILA_ENCABEZADO, column=len(ENCABEZADOS) + 1, value=COLUMNA_CONTROL)
    for fila in range(FILA_DATOS_DESDE, FILA_DATOS_HASTA + 1):
        obs = hoja.cell(row=fila, column=14)
        obs.value = None
        control = hoja.cell(row=fila, column=15)
        control.value = _formula_control(fila)
    if "Instrucciones" not in wb.sheetnames:
        hoja_inst = wb.create_sheet("Instrucciones")
    else:
        hoja_inst = wb["Instrucciones"]
    hoja_inst.delete_rows(1, hoja_inst.max_row or 1)
    for i, linea in enumerate(INSTRUCCIONES, start=1):
        hoja_inst.cell(row=i, column=1, value=linea)
    wb.save(ruta)


def main() -> int:
    for destino in DESTINOS:
        if not destino.is_file():
            raise SystemExit(f"No existe {destino}")
        actualizar(destino)
        print(f"Actualizado {destino}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
